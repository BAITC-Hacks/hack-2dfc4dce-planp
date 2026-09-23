"""Integration checks use only supplied workbooks, never invented business rows."""
import os
from pathlib import Path

from openpyxl import load_workbook
import pytest

from smartbuyer.data import FILES, load_dataset


DATA_ROOT = Path(os.environ.get("HACKALEM_DATA_DIR", os.environ.get("HACKALEM_DATA_ROOT", str(
    Path.home() / "OneDrive/Desktop/HackAlem_Logistics_2026-09-23/04_Excel"))))


@pytest.fixture(scope="module")
def dataset():
    if not DATA_ROOT.is_dir():
        pytest.skip("Set HACKALEM_DATA_DIR to the supplied 04_Excel directory.")
    return load_dataset(DATA_ROOT)


def rows(supplier, role, sheet):
    book = load_workbook(DATA_ROOT / FILES[supplier][role], read_only=True, data_only=True, keep_links=False)
    try:
        return list(book[sheet].iter_rows(values_only=True))
    finally:
        book.close()


def test_unique_keys_and_all_provided_files(dataset):
    keys = [(i["supplier"], i["sku"]) for i in dataset["items"]]
    assert len(keys) == len(set(keys))
    assert len(dataset["sources"]) == 12
    assert all(len(s["sha256"]) == 64 for s in dataset["sources"])
    assert dataset["partial_months"] == ["2026-09"]
    assert dataset["as_of"] == "2026-09-22"
    assert "filename" in dataset["as_of_basis"]


@pytest.mark.parametrize("supplier,key_column,start_column", [("SystemElectric", 2, 4), ("IEK", 1, 2)])
def test_monthly_rows_and_quantities_match_independent_read(dataset, supplier, key_column, start_column):
    original = rows(supplier, "sales", "Лист_1")
    expected = {str(r[key_column]).strip(): r for r in original[2:] if r[key_column] is not None}
    actual = {i["sku"]: i for i in dataset["items"] if i["supplier"] == supplier}
    assert set(expected) <= set(actual)
    for sku, row in expected.items():
        item = actual[sku]
        assert len(item["history"]) == 33
        assert list(item["history"])[0] == "2024-01"
        assert list(item["history"])[-1] == "2026-09"
        assert item["partial_months"] == ["2026-09"]
        assert list(item["history"].values()) == list(row[start_column:start_column + 33])
    assert any(v is None for item in actual.values() for v in item["history"].values())
    assert any(v is not None and v < 0 for item in actual.values() for v in item["history"].values())


def test_snapshot_join_reserve_once_and_no_guessed_eta(dataset):
    original = rows("SystemElectric", "snapshot", "TDSheet")
    actual = {i["sku"]: i for i in dataset["items"] if i["supplier"] == "SystemElectric"}
    for row in original[2:]:
        item = actual[str(row[1]).strip()]
        assert item["stock"]["reported"] == row[49]
        assert item["stock"]["reserved"] == row[50]
        assert item["stock"]["free"] == row[51]
        assert item["stock"]["free"] == item["stock"]["reported"] - item["stock"]["reserved"]
        assert item["stock"]["status"] == "reported_not_confirmed"
        assert item["inbound"][0]["quantity"] == row[54]
        assert item["inbound"][0]["eta"] is None
    assert sum(i["stock"]["status"] == "reported_not_confirmed" for i in actual.values()) == len(original) - 2


def test_pack_by_article_and_missing_not_one(dataset):
    original = rows("SystemElectric", "pack", "Лист_1")
    expected = {str(r[3]).strip(): r[4] for r in original[2:] if r[3] is not None}
    actual = {i["sku"]: i for i in dataset["items"] if i["supplier"] == "SystemElectric"}
    assert all(i["pack_multiple"] == expected.get(sku) for sku, i in actual.items())
    assert any(i["pack_multiple"] is None and i["stock"]["reported"] is not None for i in actual.values())
    assert all(i["moq"] is None for i in actual.values())


def test_iek_month_start_stock_not_current(dataset):
    items = [i for i in dataset["items"] if i["supplier"] == "IEK"]
    assert items
    assert all(i["stock"]["reported"] is None and i["stock"]["free"] is None for i in items)
    assert all(i["pack_multiple"] is None for i in items)
    assert any(i["inbound"] for i in items)
    assert all(p["status"] == "reported_deadline_not_confirmed" for i in items for p in i["inbound"])


@pytest.mark.parametrize("supplier,unit_column", [("SystemElectric", 3), ("IEK", 1)])
def test_units_join_by_internal_code_and_reference(dataset, supplier, unit_column):
    original = rows(supplier, "stock_history", "Лист_1")
    expected = {str(row[2]).strip(): (number, row[unit_column])
                for number, row in enumerate(original[2:], 3) if row[2] is not None}
    matched = 0
    for item in (i for i in dataset["items"] if i["supplier"] == supplier):
        source = expected.get(item["internal_code"])
        if source is None:
            assert item["unit"] is None
            continue
        matched += 1
        number, value = source
        assert item["unit"] == (str(value).strip() or None if value is not None else None)
        assert {"file": FILES[supplier]["stock_history"], "sheet": "Лист_1",
                "range": f"{'D' if unit_column == 3 else 'B'}{number}", "field": "unit"} in item["source_refs"]
    assert matched > 0


def test_missing_root_has_clear_error(tmp_path):
    with pytest.raises(ValueError, match="Не найден исходный XLSX"):
        load_dataset(tmp_path)
