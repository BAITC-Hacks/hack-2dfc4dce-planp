"""CLI integration against the issued files; no generated business dataset."""

from calendar import monthrange
import csv
import io
import json
import os
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest

from smartbuyer.cli import _csv_cell, build_report, render_report
from smartbuyer.data import FILES, load_dataset
from smartbuyer.forecast import forecast_month


DATA_ROOT = Path(os.environ.get("HACKALEM_DATA_DIR", os.environ.get("HACKALEM_DATA_ROOT",
    str(Path.home() / "OneDrive/Desktop/HackAlem_Logistics_2026-09-23/04_Excel"))))


@pytest.fixture(scope="module")
def dataset():
    if not DATA_ROOT.is_dir():
        pytest.skip("Provide issued workbooks via HACKALEM_DATA_DIR; no substitute data is generated.")
    return load_dataset(DATA_ROOT)


@pytest.fixture(scope="module")
def report(dataset):
    return build_report(dataset, "2026-10", "ATN544045")


def test_known_sku_forecasts_match_source_arithmetic(report):
    book = load_workbook(DATA_ROOT / FILES["SystemElectric"]["sales"], read_only=True, data_only=True)
    try:
        source = next(row for row in book["Лист_1"].iter_rows(values_only=True) if row[2] == "ATN544045")
    finally:
        book.close()
    history = dict(zip((f"{2024+i//12}-{i%12+1:02d}" for i in range(33)), source[4:37]))
    row = report["rows"][0]
    growth = sum(history[f"2026-{m:02d}"] for m in (6, 7, 8)) / sum(history[f"2025-{m:02d}"] for m in (6, 7, 8))
    months = [f"2025-{m:02d}" for m in range(9, 13)] + [f"2026-{m:02d}" for m in range(1, 9)]
    daily_average = sum(history[m] / monthrange(int(m[:4]), int(m[5:]))[1] for m in months) / 12
    assert row["forecasts"]["seasonal_growth"]["forecast_units"] == pytest.approx(history["2025-10"] * growth)
    assert row["forecasts"]["mean_12"]["forecast_units"] == pytest.approx(daily_average * 31)
    assert report["dataset_as_of"] == report["as_of"] == dataset_date_from_filename()
    assert report["sources"] and row["source_refs"] and row["unit"]


def dataset_date_from_filename():
    name = FILES["SystemElectric"]["snapshot"]
    day, month, year = name.rsplit(" на ", 1)[1].removesuffix(".xlsx").split(".")
    return f"{year}-{month}-{day}"


def test_json_csv_share_results_statuses_nulls_and_sources(report):
    parsed = json.loads(render_report(report))
    csv_row = next(csv.DictReader(io.StringIO(render_report(parsed, "csv"))))
    row = parsed["rows"][0]
    assert float(csv_row["seasonal_growth_units"]) == row["forecasts"]["seasonal_growth"]["forecast_units"]
    assert float(csv_row["mean_12_units"]) == row["forecasts"]["mean_12"]["forecast_units"]
    assert float(csv_row["seasonal_mae_units"]) == row["backtest"]["metrics"]["seasonal_growth"]["mae_units"]
    assert csv_row["calculation_status"] == "ok"
    assert csv_row["forecast_status"] == "provisional"
    assert csv_row["order_status"] == "blocked" and csv_row["order_quantity"] == ""
    assert row["order_quantity"] is None and json.loads(csv_row["order_reasons"]) == row["order_reasons"]
    assert json.loads(csv_row["source_refs"]) == row["source_refs"]
    assert json.loads(csv_row["sources"]) == report["sources"]
    assert csv_row["stock_status"] == row["stock"]["status"]


def test_backtest_complete_past_months_and_metrics(dataset, report):
    row = report["rows"][0]
    item = next(item for item in dataset["items"] if item["sku"] == row["sku"])
    assert row["backtest_months"] == [f"2026-{m:02d}" for m in range(3, 9)]
    assert not set(row["backtest_months"]) & set(item["partial_months"])
    assert row["backtest"]["tested"] == len(row["backtest_months"])
    for point in row["backtest"]["points"]:
        past = {month: value for month, value in item["history"].items() if month < point["month"]}
        for method, quantity in point["forecasts"].items():
            expected = forecast_month(past, point["month"], point["month"] + "-01", method)
            assert quantity == expected["forecast_units"]


def test_reported_free_stock_is_preserved_once(dataset):
    item = next(item for item in dataset["items"] if (item["stock"]["reserved"] or 0) > 0)
    row = next(row for row in build_report(dataset, "2026-10", item["sku"])["rows"] if row["supplier"] == item["supplier"])
    assert row["stock"] == item["stock"]
    assert row["stock"]["free"] == row["stock"]["reported"] - row["stock"]["reserved"]
    assert row["order_quantity"] is None


def test_missing_month_is_blocked_not_zero_and_no_mixed_mae(dataset):
    item = next(item for item in dataset["items"] if item["history"].get("2025-10") is None)
    report = build_report(dataset, "2026-10", item["sku"])
    row = next(row for row in report["rows"] if row["supplier"] == item["supplier"])
    assert row["forecast_status"] == "blocked"
    assert row["forecasts"]["seasonal_growth"]["forecast_units"] is None
    assert "нет данных" in row["explanation"] and "None" not in row["explanation"]
    assert not any("mae" in key for key in report["coverage"])


def test_future_override_rejected(dataset):
    with pytest.raises(ValueError, match="позже"):
        build_report(dataset, "2026-10", as_of="2026-09-23")


@pytest.mark.parametrize("text", ["=1+1", "+SUM(A1)", "-cmd", "@SUM(A1)", " \t=1", "\r=1"])
def test_csv_formula_guard_for_untrusted_text(text):
    # Security strings are not invented business observations or a demo dataset.
    assert _csv_cell(text) == "'" + text


def test_live_cli_utf8_json_and_no_traceback(dataset, tmp_path):
    command = [sys.executable, "-m", "smartbuyer.cli", "--target", "2026-10", "--sku", "ATN544045"]
    success = subprocess.run(command + ["--data-dir", str(DATA_ROOT)], capture_output=True, encoding="utf-8", check=False)
    assert success.returncode == 0 and success.stderr == ""
    assert json.loads(success.stdout)["rows"][0]["sku"] == "ATN544045"
    missing = subprocess.run(command + ["--data-dir", str(tmp_path / "absent")], capture_output=True, encoding="utf-8", check=False)
    assert missing.returncode == 2 and missing.stdout == ""
    assert "Ошибка данных" in missing.stderr and "Traceback" not in missing.stderr
