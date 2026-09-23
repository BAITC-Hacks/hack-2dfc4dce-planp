"""Read the supplied HackAlem workbooks without executing or changing formulas."""

from collections import Counter
from datetime import datetime
from hashlib import sha256
from math import isfinite, isclose
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


AS_OF = "2026-09-22"
AS_OF_BASIS = "filename, not verified live inventory"
FILES = {
    "SystemElectric": {
        "sales": "Systeme electric/Ежемесячные продажи в кол-м выражении SystemElectric 2024-2026.xlsx",
        "snapshot": "Systeme electric/Товар в пути_SystemElectric на 22.09.2026.xlsx",
        "pack": "Systeme electric/MOQ SystemElectric.xlsx",
        "stock_history": "Systeme electric/Ежемесячные остатки SystemElectric 2024-2026.xlsx",
        "transactions": "Systeme electric/Динамика продаж_Syseme Electric_2025-2026.xlsx",
        "seasonality": "Systeme electric/Сезонность SystemElectric 2024-2026.xlsx",
    },
    "IEK": {
        "sales": "IEK/Ежемесячные продажи в количественном выражении за последние 2 года.xlsx",
        "snapshot": "IEK/Путь ИЭК 22.09.2026.xlsx",
        "pack": "IEK/MOQ  ИЭК.xlsx",
        "stock_history": "IEK/Ежемесячные остатки продукции за последние 2 года  ИЭК.xlsx",
        "transactions": "IEK/Динамика продаж_2025-2026.xlsx",
        "seasonality": "IEK/Сезонность ИЭК.xlsx",
    },
}
_MONTHS = {v: i + 1 for i, v in enumerate(
    ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")
)}


def _text(value):
    return "" if value is None else str(value).strip()


def _number(value, warnings, where):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
        warnings.append(f"{where}: нечисловое значение/ошибка Excel, требуется проверка.")
        return None
    value = float(value)
    if value < 0:
        warnings.append(f"{where}: отрицательное значение сохранено; смысл возврата/корректировки не подтверждён.")
    return value


def _read(root, relative, sheet, sources, supplier, role):
    path = (root / relative).resolve()
    if not path.is_relative_to(root):
        raise ValueError(f"Источник выходит за data_root: {relative}")
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise ValueError(f"Не найден исходный XLSX: {relative}")
    try:
        with ZipFile(path) as archive:
            if any("vbaproject" in n.lower() or n.startswith("xl/activeX/") for n in archive.namelist()):
                raise ValueError(f"Активное содержимое запрещено: {relative}")
        source = {"file": relative, "supplier": supplier, "role": role,
                  "sha256": sha256(path.read_bytes()).hexdigest(),
                  "status": "not_parsed" if sheet is None else "read_cached_values"}
        sources.append(source)
        if sheet is None:
            return []
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if sheet not in book.sheetnames:
                raise ValueError(f"Лист {sheet} отсутствует: {relative}")
            return list(book[sheet].iter_rows(values_only=True))
        finally:
            book.close()
    except (BadZipFile, OSError) as exc:
        raise ValueError(f"Не удалось прочитать XLSX: {relative}") from exc


def _index(rows, key_column, start, warnings, source):
    """Ambiguous keys are unavailable, never silently first-match or summed."""
    counts = Counter(_text(r[key_column]) for r in rows[start:] if _text(r[key_column]))
    duplicates = {key for key, count in counts.items() if count > 1}
    if duplicates:
        warnings.append(f"{source}: неоднозначных ключей {len(duplicates)}; их значения не объединяются.")
    return {_text(r[key_column]): (n, r) for n, r in enumerate(rows[start:], start + 1)
            if _text(r[key_column]) and _text(r[key_column]) not in duplicates}, duplicates


def _months(header):
    result = {}
    for index, value in enumerate(header):
        if isinstance(value, datetime):
            month = value.strftime("%Y-%m")
        else:
            match = re.fullmatch(r"([А-Яа-яЁё]+)\.?\s+(\d{4})", _text(value))
            if not match or match[1].lower()[:3] not in _MONTHS:
                continue
            month = f"{match[2]}-{_MONTHS[match[1].lower()[:3]]:02d}"
        if month in result.values():
            raise ValueError(f"Дублирующийся месяц в заголовке: {month}")
        result[index] = month
    if not result:
        raise ValueError("В таблице продаж не найдены месяцы.")
    return result


def _check_header(rows, row_number, expected, source):
    if len(rows) < row_number:
        raise ValueError(f"Нет заголовка таблицы: {source}")
    header = rows[row_number - 1]
    for column, label in expected.items():
        if column >= len(header) or _text(header[column]) != label:
            raise ValueError(f"Изменилась схема {source}: ожидался {label} в {get_column_letter(column + 1)}{row_number}.")


def _item(sku, supplier, name):
    return {"sku": sku, "supplier": supplier, "name": _text(name), "unit": None,
            "category": None, "history": {}, "partial_months": [],
            "stock": {"reported": None, "reserved": None, "free": None,
                      "as_of": None, "status": "missing"},
            "pack_multiple": None, "moq": None, "inbound": [], "warnings": [], "source_refs": []}


def _ref(item, relative, sheet, cell_range, field):
    item["source_refs"].append({"file": relative, "sheet": sheet, "range": cell_range, "field": field})


def load_dataset(data_root: str | Path) -> dict:
    """Load supplied tables; missing values and unverified stock remain explicit.

    SystemElectric joins supplier articles. IEK sales lack articles, so its SKU
    is the internal 1C code and supplier_article is attached only when unambiguous.
    Monthly sales are an explicit provisional source, not reconciled transactions.
    """
    root = Path(data_root).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("data_root должен указывать на каталог с выданными XLSX.")
    result = {"as_of": AS_OF, "as_of_basis": AS_OF_BASIS, "items": [], "sources": [],
              "partial_months": [AS_OF[:7]], "warnings": [
                  "База спроса — ежемесячные продажи, временное допущение до сверки с операциями.",
                  "2026-09 — неполный месяц; не использовать как завершённый период.",
                  "Пустые ячейки неизвестны, не нули. Отрицательные продажи требуют разбора.",
                  "Прочитаны сохранённые значения Excel; формулы не исполнялись и не пересчитывались.",
                  "Дата 2026-09-22 взята из имён файлов, это не подтверждение живых остатков.",
                  "Нет подтверждённых клиентских заказов, дневной доступности и срока новой поставки.",
                  "Таблицы операций и готовых коэффициентов не смешиваются с месячными продажами.",
              ]}
    for supplier, files in FILES.items():
        systeme = supplier == "SystemElectric"
        sheets = {"sales": "Лист_1", "snapshot": "TDSheet" if systeme else "Лист4",
                  "pack": "Лист_1" if systeme else "Лист7", "stock_history": "Лист_1"}
        tables = {role: _read(root, relative, sheets.get(role), result["sources"], supplier, role)
                  for role, relative in files.items()}
        _check_header(tables["sales"], 1, {0: "Номенклатура", 1: "Номенклатура.Код"}, files["sales"])
        _check_header(tables["pack"], 1, {4: "Кратность" if systeme else "Мин. разр. к отгр."}, files["pack"])
        _check_header(tables["stock_history"], 1, {2: "Номенклатура.Код", 3 if systeme else 1: "Ед.изм" if systeme else "Ед."}, files["stock_history"])
        if systeme:
            _check_header(tables["sales"], 1, {2: "Артикул"}, files["sales"])
            _check_header(tables["pack"], 1, {3: "Артикул"}, files["pack"])
            _check_header(tables["snapshot"], 2, {1: "Артикул поставщика", 49: "Остаток", 50: "Зарезервировано", 51: "Свободный остаток", 54: "СЭ в пути 24.09"}, files["snapshot"])
        else:
            _check_header(tables["pack"], 1, {1: "Код 1с", 2: "Артикул поставщика"}, files["pack"])
            _check_header(tables["snapshot"], 1, {0: "Код 1с", 1: "Артикул ИЭК"}, files["snapshot"])
        sales = tables["sales"]
        months = _months(sales[0])
        sales_index, sales_dupes = _index(sales, 2 if systeme else 1, 2, result["warnings"], files["sales"])
        snapshot_index, snapshot_dupes = _index(tables["snapshot"], 1 if systeme else 0,
                                               2 if systeme else 1, result["warnings"], files["snapshot"])
        pack_index, pack_dupes = _index(tables["pack"], 3 if systeme else 1,
                                       2 if systeme else 1, result["warnings"], files["pack"])
        unit_index, unit_dupes = _index(tables["stock_history"], 2, 2, result["warnings"], files["stock_history"])
        # ponytail: fixed supplied layouts, not a general spreadsheet-mapping engine.
        keys = set(sales_index) | sales_dupes
        if systeme:
            keys |= set(snapshot_index) | snapshot_dupes
        for sku in sorted(keys):
            sales_row = sales_index.get(sku)
            snapshot_row = snapshot_index.get(sku)
            name = sales_row[1][0] if sales_row else (snapshot_row[1][3] if snapshot_row else sku)
            item = _item(sku, supplier, name)
            item["sku_kind"] = "supplier_article" if systeme else "internal_code"
            code = _text(sales_row[1][1]) if sales_row else (_text(snapshot_row[1][2]) if snapshot_row else "")
            item["internal_code"] = code
            item["supplier_article"] = sku if systeme else None
            if sales_row:
                n, row = sales_row
                item["history"] = {month: _number(row[col], item["warnings"], f"sales!{get_column_letter(col+1)}{n}")
                                   for col, month in months.items()}
                item["partial_months"] = [m for m in item["history"] if m == AS_OF[:7]]
                _ref(item, files["sales"], "Лист_1", f"A{n}:{get_column_letter(max(months)+1)}{n}", "history")
            else:
                item["warnings"].append("Нет однозначной строки в выбранной базе месячных продаж.")
            unit_row = unit_index.get(code)
            if unit_row:
                n, row = unit_row
                item["unit"] = _text(row[3 if systeme else 1]) or None
                _ref(item, files["stock_history"], "Лист_1", f"{'D' if systeme else 'B'}{n}", "unit")
            if code in unit_dupes:
                item["warnings"].append("Неоднозначная единица измерения по коду.")
            pack_row = pack_index.get(sku)
            if pack_row:
                n, row = pack_row
                value = _number(row[4], item["warnings"], f"MOQ!E{n}")
                if value is not None and value <= 0:
                    item["warnings"].append("MOQ/кратность должны быть положительными; значение не применяется.")
                    value = None
                field = "pack_multiple" if systeme else "moq"
                item[field] = value
                if not systeme:
                    item["supplier_article"] = _text(row[2]) or None
                _ref(item, files["pack"], sheets["pack"], f"E{n}", field)
            if sku in pack_dupes:
                item["warnings"].append("Дублирующийся ключ MOQ/кратности; значение не применяется.")
            if item["pack_multiple"] is None:
                item["warnings"].append("Кратность поставки неизвестна; не подставлять 1.")
            if systeme:
                item["warnings"].append("MOQ не задан отдельно от кратности; не подменять эти понятия.")
                if snapshot_row:
                    n, row = snapshot_row
                    item["category"] = _text(row[4]) or None
                    item["stock"] = {field: _number(row[col], item["warnings"], f"TDSheet!{get_column_letter(col+1)}{n}")
                                     for field, col in (("reported", 49), ("reserved", 50), ("free", 51))}
                    item["stock"].update(as_of=AS_OF, status="reported_not_confirmed")
                    stock = item["stock"]
                    if all(stock[f] is not None for f in ("reported", "reserved", "free")):
                        if not isclose(stock["reported"] - stock["reserved"], stock["free"], abs_tol=1e-9):
                            stock["status"] = "inconsistent_reported"
                            item["warnings"].append("AX − AY не равно AZ: требуется сверка остатков.")
                    item["inbound"] = [{"quantity": _number(row[54], item["warnings"], f"TDSheet!BC{n}"),
                                        "eta": None, "eta_label": _text(tables["snapshot"][1][54]),
                                        "status": "reported_eta_year_unknown"}]
                    _ref(item, files["snapshot"], "TDSheet", f"AX{n}:AZ{n}", "stock")
                    _ref(item, files["snapshot"], "TDSheet", f"BC{n}", "inbound")
                    _ref(item, files["snapshot"], "TDSheet", f"E{n}", "category")
                    item["warnings"].append("В ETA «24.09» не указан год; поставка не подтверждена.")
                    item["warnings"].append("Код категории сохранён без выдуманной расшифровки.")
                else:
                    item["warnings"].append("Нет однозначного текущего отчётного остатка.")
            else:
                item["warnings"].append("Исторические остатки на начало месяца не являются текущим складом.")
                if snapshot_row:
                    n, row = snapshot_row
                    item["supplier_article"] = item["supplier_article"] or _text(row[1]) or None
                    for col in range(3, 9):
                        if row[col] is None:
                            continue
                        label = _text(tables["snapshot"][0][col])
                        match = re.search(r"поступление до (\d{2}\.\d{2}\.\d{4})", label)
                        eta = datetime.strptime(match[1], "%d.%m.%Y").date().isoformat() if match else None
                        item["inbound"].append({"quantity": _number(row[col], item["warnings"], f"Лист4!{get_column_letter(col+1)}{n}"),
                                                "eta": eta, "eta_label": label, "status": "reported_deadline_not_confirmed"})
                    _ref(item, files["snapshot"], "Лист4", f"D{n}:I{n}", "inbound")
                item["warnings"].append("Отсутствующие строки/ячейки пути не доказывают отсутствие поставок.")
            if sku in snapshot_dupes:
                item["warnings"].append("Дублирующийся ключ в пути/остатках: связанные значения недоступны.")
            if any(v is None for v in item["history"].values()):
                item["warnings"].append("Есть неизвестные месяцы продаж: пустые ячейки не приравнены нулю.")
            result["items"].append(item)
    return result
